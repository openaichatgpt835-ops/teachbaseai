"""KB ingestion: parse files, chunk text, generate embeddings."""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import subprocess
import tempfile
import base64
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from apps.backend.models.kb import KBFile, KBChunk, KBEmbedding
from apps.backend.services.kb_settings import get_effective_gigachat_settings, get_valid_gigachat_access_token
from apps.backend.services.gigachat_client import create_embeddings
from apps.backend.services.kb_pgvector import write_vector_column
from apps.backend.services.billing import get_pricing, calc_cost_rub, record_usage


def _read_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _read_csv_file(path: str) -> str:
    lines: list[str] = []
    with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row:
                continue
            lines.append(" | ".join(str(c).strip() for c in row if c is not None))
    return "\n".join(lines)


def _read_xlsx_file(path: str) -> str:
    import openpyxl  # type: ignore
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                line = " | ".join(str(c).strip() for c in row if c is not None)
                if line:
                    lines.append(line)
    finally:
        wb.close()
    return "\n".join(lines)


def _read_xls_file(path: str) -> str:
    import xlrd  # type: ignore
    book = xlrd.open_workbook(path)
    lines: list[str] = []
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            row = sheet.row_values(r)
            line = " | ".join(str(c).strip() for c in row if c is not None and str(c).strip())
            if line:
                lines.append(line)
    return "\n".join(lines)


def _read_docx_file(path: str) -> str:
    import docx  # type: ignore
    d = docx.Document(path)
    parts: list[str] = []
    for p in d.paragraphs:
        if p.text and p.text.strip():
            parts.append(p.text.strip())
    return "\n".join(parts)


def _read_pdf_file(path: str) -> str:
    from pypdf import PdfReader  # type: ignore
    reader = PdfReader(path)
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    return "\n".join(parts)


def _read_pdf_pages(path: str) -> list[tuple[int, str]]:
    from pypdf import PdfReader  # type: ignore
    reader = PdfReader(path)
    out: list[tuple[int, str]] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        out.append((idx, text.strip()))
    return out


def _read_epub_file(path: str) -> str:
    from ebooklib import epub, ITEM_DOCUMENT  # type: ignore
    from bs4 import BeautifulSoup  # type: ignore
    book = epub.read_epub(path)
    parts: list[str] = []
    for item in book.get_items():
        item_type = item.get_type() if hasattr(item, "get_type") else None
        media_type = getattr(item, "media_type", "") or ""
        is_doc = (item_type == ITEM_DOCUMENT) or media_type in ("application/xhtml+xml", "text/html")
        if is_doc:
            soup = BeautifulSoup(item.get_body_content(), "lxml")
            text = soup.get_text("\n").strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _read_fb2_file(path: str) -> str:
    from bs4 import BeautifulSoup  # type: ignore
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        data = f.read()
    soup = BeautifulSoup(data, "lxml-xml")
    texts = []
    for p in soup.find_all(["p", "subtitle", "title"]):
        t = (p.get_text() or "").strip()
        if t:
            texts.append(t)
    return "\n".join(texts)


def _ocr_pdf_file(path: str, lang: str = "rus+eng") -> str:
    from pdf2image import convert_from_path  # type: ignore
    import pytesseract  # type: ignore
    pages = convert_from_path(path, dpi=200)
    parts: list[str] = []
    for img in pages:
        text = pytesseract.image_to_string(img, lang=lang) or ""
        if text.strip():
            parts.append(text.strip())
    return "\n".join(parts)


def _ocr_pdf_pages(path: str, lang: str = "rus+eng") -> list[tuple[int, str]]:
    from pdf2image import convert_from_path  # type: ignore
    import pytesseract  # type: ignore
    pages = convert_from_path(path, dpi=200)
    out: list[tuple[int, str]] = []
    for idx, img in enumerate(pages, start=1):
        text = (pytesseract.image_to_string(img, lang=lang) or "").strip()
        out.append((idx, text))
    return out


def _preview_pdf_path(storage_path: str) -> str:
    return f"{storage_path}.preview.pdf"


def _build_epub_preview_html(path: str) -> str | None:
    from ebooklib import epub, ITEM_DOCUMENT  # type: ignore
    from bs4 import BeautifulSoup  # type: ignore
    book = epub.read_epub(path)
    id_to_item: dict[str, object] = {}
    image_map: dict[str, str] = {}
    for item in book.get_items():
        item_id = getattr(item, "id", None)
        if item_id:
            id_to_item[str(item_id)] = item
        media_type = (getattr(item, "media_type", "") or "").strip().lower()
        if media_type.startswith("image/"):
            raw = getattr(item, "get_content", lambda: b"")() or b""
            if not raw:
                continue
            b64 = base64.b64encode(raw).decode("ascii")
            data_uri = f"data:{media_type};base64,{b64}"
            fname = (getattr(item, "file_name", "") or "").strip()
            href = (getattr(item, "href", "") or "").strip()
            for key in (fname, href, fname.lstrip("./"), href.lstrip("./"), os.path.basename(fname), os.path.basename(href)):
                if key:
                    image_map[key] = data_uri

    ordered_docs: list[object] = []
    for sp in (getattr(book, "spine", None) or []):
        sid = sp[0] if isinstance(sp, (list, tuple)) and sp else sp
        it = id_to_item.get(str(sid))
        if not it:
            continue
        media_type = (getattr(it, "media_type", "") or "").strip().lower()
        if media_type in ("application/xhtml+xml", "text/html"):
            ordered_docs.append(it)
    if not ordered_docs:
        for item in book.get_items():
            media_type = (getattr(item, "media_type", "") or "").strip().lower()
            item_type = item.get_type() if hasattr(item, "get_type") else None
            if item_type == ITEM_DOCUMENT or media_type in ("application/xhtml+xml", "text/html"):
                ordered_docs.append(item)

    parts: list[str] = []
    for item in ordered_docs:
        soup = BeautifulSoup(item.get_body_content(), "lxml")
        for bad in soup.find_all(["script", "style"]):
            bad.decompose()
        for img in soup.find_all("img"):
            src = (img.get("src") or "").strip()
            if not src:
                continue
            src_clean = src.split("#", 1)[0]
            candidates = [src_clean, src_clean.lstrip("./"), os.path.basename(src_clean)]
            repl = next((image_map.get(c) for c in candidates if c in image_map), None)
            if repl:
                img["src"] = repl
                img["style"] = "max-width:100%;height:auto;display:block;margin:8pt auto;"
        body = soup.body or soup
        chapter = str(body)
        if chapter.strip():
            parts.append(f"<section class='chapter'>{chapter}</section>")
    if not parts:
        return None
    return (
        "<html><head><meta charset='utf-8'/>"
        "<style>"
        "body{font-family:Arial,sans-serif;font-size:11pt;line-height:1.45;margin:16mm;text-align:left;}"
        ".chapter{page-break-after:always;}"
        "p{margin:0 0 8pt 0; text-align:left;}"
        "h1,h2,h3,h4{margin:10pt 0 8pt 0;}"
        "img{max-width:100%;height:auto;display:block;margin:8pt auto;}"
        "</style></head><body>"
        + "\n".join(parts)
        + "</body></html>"
    )


def _build_fb2_preview_html(path: str) -> str | None:
    from bs4 import BeautifulSoup, NavigableString  # type: ignore
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        data = f.read()
    soup = BeautifulSoup(data, "lxml-xml")
    binaries: dict[str, str] = {}
    for b in soup.find_all("binary"):
        bid = (b.get("id") or "").strip()
        ctype = (b.get("content-type") or "image/jpeg").strip()
        raw_b64 = "".join((b.get_text() or "").split())
        if bid and raw_b64:
            binaries[bid] = f"data:{ctype};base64,{raw_b64}"

    body = soup.find("body")
    if not body:
        return None

    def render_node(node) -> str:
        if isinstance(node, NavigableString):
            txt = str(node)
            return txt if txt.strip() else ""
        if not getattr(node, "name", None):
            return ""
        name = str(node.name).lower()
        if name == "image":
            href = (node.get("l:href") or node.get("xlink:href") or node.get("href") or "").strip()
            bid = href[1:] if href.startswith("#") else href
            src = binaries.get(bid, "")
            if not src:
                return ""
            return f"<img src='{src}' style='max-width:100%;height:auto;display:block;margin:8pt auto;'/>"
        children = "".join(render_node(ch) for ch in getattr(node, "children", [])).strip()
        if not children:
            return ""
        if name in ("title",):
            return f"<h2>{children}</h2>"
        if name in ("subtitle",):
            return f"<h3>{children}</h3>"
        if name in ("section",):
            return f"<section class='chapter'>{children}</section>"
        if name in ("p", "v", "text-author", "epigraph", "cite", "poem", "stanza"):
            return f"<p>{children}</p>"
        return children

    html_body = "".join(render_node(ch) for ch in getattr(body, "children", []))
    if not html_body.strip():
        return None
    return (
        "<html><head><meta charset='utf-8'/>"
        "<style>"
        "body{font-family:Arial,sans-serif;font-size:11pt;line-height:1.45;margin:16mm;text-align:left;}"
        ".chapter{page-break-after:always;}"
        "p{margin:0 0 8pt 0; text-align:left;}"
        "h1,h2,h3,h4{margin:10pt 0 8pt 0;}"
        "img{max-width:100%;height:auto;display:block;margin:8pt auto;}"
        "</style></head><body>"
        + html_body
        + "</body></html>"
    )


def _generate_preview_pdf(src_path: str) -> str | None:
    """Best-effort conversion to PDF via LibreOffice (soffice)."""
    def _convert_with_soffice(input_path: str) -> str | None:
        with tempfile.TemporaryDirectory() as td:
            cmd = [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                td,
                input_path,
            ]
            subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            base = os.path.splitext(os.path.basename(input_path))[0]
            candidate = os.path.join(td, f"{base}.pdf")
            if not os.path.exists(candidate):
                return None
            out_path = _preview_pdf_path(src_path)
            with open(candidate, "rb") as rf, open(out_path, "wb") as wf:
                wf.write(rf.read())
            return out_path

    def _fallback_book_to_html_pdf() -> str | None:
        ext = os.path.splitext(src_path)[1].lower()
        if ext not in (".epub", ".fb2"):
            return None
        try:
            html = _build_epub_preview_html(src_path) if ext == ".epub" else _build_fb2_preview_html(src_path)
            html = (html or "").strip()
            if not html:
                return None
            with tempfile.TemporaryDirectory() as td:
                html_path = os.path.join(td, "book_preview.html")
                with open(html_path, "w", encoding="utf-8") as f:
                    f.write(html)
                return _convert_with_soffice(html_path)
        except Exception:
            return None

    try:
        out = _convert_with_soffice(src_path)
        if out:
            return out
        return _fallback_book_to_html_pdf()
    except Exception:
        return None


def _best_page_for_text(text: str, pages: list[tuple[int, str]]) -> int | None:
    t = (text or "").lower()
    if not t:
        return None
    toks = re.findall(r"[a-zа-яё0-9]{3,}", t, flags=re.IGNORECASE)
    if not toks:
        return None
    top = toks[:24]
    best: tuple[int, int] | None = None  # score, page_num
    for pnum, ptext in pages:
        plow = (ptext or "").lower()
        if not plow:
            continue
        score = sum(1 for tok in top if tok in plow)
        if score <= 0:
            continue
        if best is None or score > best[0]:
            best = (score, int(pnum))
    return best[1] if best else None


def _assign_chunk_pages_from_preview(chunks: list[KBChunk], preview_pdf_path: str) -> None:
    try:
        pages = _read_pdf_pages(preview_pdf_path)
    except Exception:
        return
    if not pages:
        return
    for ch in chunks:
        if ch.page_num is not None and int(ch.page_num) > 0:
            continue
        page = _best_page_for_text(ch.text or "", pages)
        if page:
            ch.page_num = int(page)


_VIDEO_EXTS = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"}
_MEDIA_CHUNK_SECONDS = max(60, int(os.getenv("MEDIA_CHUNK_SECONDS") or 600))
_MEDIA_CHUNK_OVERLAP_SECONDS = max(0, int(os.getenv("MEDIA_CHUNK_OVERLAP_SECONDS") or 2))

_CHUNK_PROFILES: dict[str, tuple[int, int]] = {
    # default narrative text
    "default": (1200, 200),
    # tables/spreadsheets/rows
    "tabular": (700, 120),
    # scanned/ocr pages benefit from shorter windows
    "ocr": (900, 140),
    # media transcript chunks
    "media": (1000, 120),
}


def _chunk_profile_for_ext(ext: str) -> tuple[int, int]:
    ext = (ext or "").lower()
    if ext in (".csv", ".xls", ".xlsx"):
        return _CHUNK_PROFILES["tabular"]
    if ext in (".png", ".jpg", ".jpeg"):
        return _CHUNK_PROFILES["ocr"]
    if ext in _VIDEO_EXTS:
        return _CHUNK_PROFILES["media"]
    return _CHUNK_PROFILES["default"]


@dataclass
class _Segment:
    text: str
    start_ms: int
    end_ms: int


def _extract_audio_to_wav(src_path: str, dst_path: str) -> None:
    cmd = [
        "ffmpeg",
        "-y",
        "-i",
        src_path,
        "-ac",
        "1",
        "-ar",
        "16000",
        dst_path,
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def _media_duration_seconds(src_path: str) -> int:
    cmd = [
        "ffprobe",
        "-v",
        "error",
        "-show_entries",
        "format=duration",
        "-of",
        "default=noprint_wrappers=1:nokey=1",
        src_path,
    ]
    out = subprocess.check_output(cmd, stderr=subprocess.DEVNULL).decode("utf-8", errors="ignore").strip()
    try:
        return max(1, int(float(out)))
    except Exception:
        return 1


_WHISPER_MODEL = None


def _is_noise_transcript_text(text: str) -> bool:
    t = (text or "").strip().lower()
    if not t:
        return True
    if len(t) < 14:
        return True
    bad_phrases = (
        "позитивная музыка",
        "подпишись",
        "ставьте лайк",
        "ставь лайк",
        "жми колокольчик",
        "ссылка в описании",
        "промокод",
        "реклама",
    )
    if any(p in t for p in bad_phrases):
        return True
    toks = re.findall(r"[a-zа-яё0-9]+", t, flags=re.IGNORECASE)
    if len(toks) < 3:
        return True
    uniq = len(set(toks))
    if uniq <= 2 and len(toks) >= 5:
        return True
    return False


def _get_whisper_model():
    global _WHISPER_MODEL
    if _WHISPER_MODEL is not None:
        return _WHISPER_MODEL
    from faster_whisper import WhisperModel  # type: ignore
    size = (os.getenv("WHISPER_MODEL_SIZE") or "small").strip()
    _WHISPER_MODEL = WhisperModel(size, device="cpu", compute_type="int8")
    return _WHISPER_MODEL


def _transcribe_media(path: str) -> list[_Segment]:
    model = _get_whisper_model()
    segments, _info = model.transcribe(path, vad_filter=True)
    out: list[_Segment] = []
    for seg in segments:
        text = (seg.text or "").strip()
        if not text or _is_noise_transcript_text(text):
            continue
        start_ms = int((seg.start or 0) * 1000)
        end_ms = int((seg.end or 0) * 1000)
        out.append(_Segment(text=text, start_ms=start_ms, end_ms=end_ms))
    return out


def _transcribe_media_window(path: str, start_sec: int, duration_sec: int) -> list[_Segment]:
    with tempfile.TemporaryDirectory() as td:
        chunk_wav = os.path.join(td, "chunk.wav")
        cmd = [
            "ffmpeg",
            "-y",
            "-ss",
            str(max(0, start_sec)),
            "-i",
            path,
            "-t",
            str(max(1, duration_sec)),
            "-ac",
            "1",
            "-ar",
            "16000",
            chunk_wav,
        ]
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return _transcribe_media(chunk_wav)


def _read_transcript_segments_jsonl(path: str) -> list[_Segment]:
    if not os.path.exists(path):
        return []
    out: list[_Segment] = []
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
                txt = str(row.get("text") or "").strip()
                if not txt:
                    continue
                out.append(
                    _Segment(
                        text=txt,
                        start_ms=int(row.get("start_ms") or 0),
                        end_ms=int(row.get("end_ms") or 0),
                    )
                )
            except Exception:
                continue
    return out


def _append_transcript_segment_jsonl(path: str, seg: _Segment) -> None:
    row = {"start_ms": seg.start_ms, "end_ms": seg.end_ms, "text": seg.text}
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(row, ensure_ascii=False) + "\n")


def _transcribe_media_resumable(
    src_path: str,
    transcript_jsonl_path: str,
    progress_cb=None,
) -> list[_Segment]:
    segments = _read_transcript_segments_jsonl(transcript_jsonl_path)
    last_end_ms = segments[-1].end_ms if segments else 0
    total_sec = _media_duration_seconds(src_path)
    overlap_ms = _MEDIA_CHUNK_OVERLAP_SECONDS * 1000
    start_sec = max(0, int(max(0, last_end_ms - overlap_ms) / 1000))
    if start_sec >= total_sec:
        if progress_cb:
            progress_cb(100)
        return segments

    pos = start_sec
    while pos < total_sec:
        win = min(_MEDIA_CHUNK_SECONDS, total_sec - pos)
        win_segments = _transcribe_media_window(src_path, pos, win)
        for seg in win_segments:
            shifted = _Segment(
                text=seg.text,
                start_ms=seg.start_ms + pos * 1000,
                end_ms=seg.end_ms + pos * 1000,
            )
            # Skip overlap duplicates on resume/chunk boundaries.
            if shifted.end_ms <= last_end_ms:
                continue
            segments.append(shifted)
            _append_transcript_segment_jsonl(transcript_jsonl_path, shifted)
            last_end_ms = max(last_end_ms, shifted.end_ms)
        pos += max(1, _MEDIA_CHUNK_SECONDS - _MEDIA_CHUNK_OVERLAP_SECONDS)
        if progress_cb and total_sec > 0:
            progress_cb(min(99, int((min(pos, total_sec) / total_sec) * 100)))

    if progress_cb:
        progress_cb(100)
    return segments


def _chunk_segments(segments: list[_Segment], max_chars: int = 1200) -> list[_Segment]:
    if not segments:
        return []
    chunks: list[_Segment] = []
    buf = ""
    start_ms = None
    end_ms = None
    for seg in segments:
        if not buf:
            buf = seg.text
            start_ms = seg.start_ms
            end_ms = seg.end_ms
            continue
        if len(buf) + 1 + len(seg.text) <= max_chars:
            buf = buf + " " + seg.text
            end_ms = seg.end_ms
            continue
        if not _is_noise_transcript_text(buf):
            chunks.append(_Segment(text=buf, start_ms=start_ms or 0, end_ms=end_ms or (start_ms or 0)))
        buf = seg.text
        start_ms = seg.start_ms
        end_ms = seg.end_ms
    if buf:
        if not _is_noise_transcript_text(buf):
            chunks.append(_Segment(text=buf, start_ms=start_ms or 0, end_ms=end_ms or (start_ms or 0)))
    return chunks


def extract_text_from_file(path: str, mime_type: str | None, filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext in (".txt", ".md"):
        return _read_text_file(path)
    if ext in (".csv",):
        return _read_csv_file(path)
    if ext in (".docx",):
        return _read_docx_file(path)
    if ext in (".pdf",):
        return _read_pdf_file(path)
    if ext in (".epub",):
        return _read_epub_file(path)
    if ext in (".fb2",):
        return _read_fb2_file(path)
    if ext in (".xlsx",):
        return _read_xlsx_file(path)
    if ext in (".xls",):
        return _read_xls_file(path)
    raise ValueError(f"unsupported_file_type:{ext or mime_type or 'unknown'}")


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 200) -> list[str]:
    cleaned = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not cleaned:
        return []
    if max_chars < 200:
        max_chars = 200
    if overlap < 0:
        overlap = 0
    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + max_chars)
        chunk = cleaned[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(cleaned):
            break
        start = max(0, end - overlap)
        if start == end:
            break
    return chunks


def chunk_text_with_page(text: str, page_num: int, max_chars: int = 1200, overlap: int = 200) -> list[tuple[str, int]]:
    return [(c, page_num) for c in chunk_text(text, max_chars=max_chars, overlap=overlap)]


def _count_tokens_approx(text: str) -> int:
    return len((text or "").split())


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def ingest_file(db: Session, file_id: int, trace_id: str | None = None) -> dict:
    rec = db.get(KBFile, file_id)
    if not rec:
        return {"ok": False, "error": "file_not_found"}
    rec.status = "processing"
    rec.error_message = None
    db.add(rec)
    db.commit()

    # reuse existing chunks if they already exist (multi-version embeddings)
    chunk_rows = db.execute(
        select(KBChunk).where(KBChunk.file_id == rec.id).order_by(KBChunk.chunk_index)
    ).scalars().all()

    if not chunk_rows:
        ext = os.path.splitext(rec.filename.lower())[1]
        max_chars, overlap = _chunk_profile_for_ext(ext)
        ocr_enabled = (os.getenv("OCR_ENABLED") or "").strip().lower() in ("1", "true", "yes", "on")
        if ext in _VIDEO_EXTS:
            try:
                transcript_jsonl_path = rec.storage_path + ".transcript.jsonl"

                def _set_progress(pct: int) -> None:
                    rec.error_message = f"transcribe_progress:{max(0, min(100, int(pct)))}"
                    db.add(rec)
                    db.commit()

                with tempfile.TemporaryDirectory() as td:
                    wav_path = os.path.join(td, "audio.wav")
                    _extract_audio_to_wav(rec.storage_path, wav_path)
                    segments = _transcribe_media_resumable(
                        wav_path,
                        transcript_jsonl_path=transcript_jsonl_path,
                        progress_cb=_set_progress,
                    )
                chunks = _chunk_segments(segments, max_chars=max_chars)
                transcript_path = rec.storage_path + ".transcript.txt"
                with open(transcript_path, "w", encoding="utf-8") as tf:
                    for seg in segments:
                        tf.write(f"[{seg.start_ms}-{seg.end_ms}] {seg.text}\n")
            except Exception as e:
                rec.status = "error"
                rec.error_message = ("transcribe_failed:" + str(e))[:200]
                db.add(rec)
                db.commit()
                return {"ok": False, "error": "transcribe_failed", "detail": rec.error_message}
        else:
            try:
                if ext == ".pdf":
                    page_texts = _read_pdf_pages(rec.storage_path)
                    ocr_by_page: dict[int, str] = {}
                    if ocr_enabled:
                        try:
                            ocr_pages = _ocr_pdf_pages(rec.storage_path)
                            ocr_by_page = {int(p): (t or "").strip() for p, t in ocr_pages}
                        except Exception:
                            ocr_by_page = {}
                    chunks = []
                    for page_num, page_text in page_texts:
                        text_for_page = (page_text or "").strip()
                        if not text_for_page and ocr_by_page:
                            text_for_page = (ocr_by_page.get(int(page_num)) or "").strip()
                        if not text_for_page:
                            continue
                        chunks.extend(chunk_text_with_page(text_for_page, page_num, max_chars=max_chars, overlap=overlap))
                    if not chunks:
                        text = ""
                    else:
                        text = None
                else:
                    text = extract_text_from_file(rec.storage_path, rec.mime_type, rec.filename)
            except Exception as e:
                rec.status = "error"
                rec.error_message = str(e)[:200]
                db.add(rec)
                db.commit()
                return {"ok": False, "error": "extract_failed", "detail": rec.error_message}
            if ext != ".pdf":
                chunks = chunk_text(text, max_chars=max_chars, overlap=overlap)

        if not chunks and ext == ".pdf":
            if ocr_enabled:
                try:
                    text = _ocr_pdf_file(rec.storage_path)
                    chunks = chunk_text(text, max_chars=max_chars, overlap=overlap)
                except Exception as e:
                    rec.status = "error"
                    rec.error_message = ("ocr_failed:" + str(e))[:200]
                    db.add(rec)
                    db.commit()
                    return {"ok": False, "error": "ocr_failed", "detail": rec.error_message}
            else:
                rec.status = "error"
                rec.error_message = "no_text_chunks (ocr_disabled)"
                db.add(rec)
                db.commit()
                return {"ok": False, "error": "no_text_chunks"}
        if not chunks:
            rec.status = "error"
            rec.error_message = "no_text_chunks_after_ocr" if ocr_enabled else "no_text_chunks"
            db.add(rec)
            db.commit()
            return {"ok": False, "error": "no_text_chunks"}

        new_rows: list[KBChunk] = []
        for idx, ch in enumerate(chunks):
            if isinstance(ch, _Segment):
                text_val = ch.text
                start_ms = ch.start_ms
                end_ms = ch.end_ms
                page_num = None
            elif isinstance(ch, tuple) and len(ch) == 2:
                text_val = ch[0]
                start_ms = None
                end_ms = None
                page_num = int(ch[1])
            else:
                text_val = ch
                start_ms = None
                end_ms = None
                page_num = None
            new_rows.append(KBChunk(
                portal_id=rec.portal_id,
                file_id=rec.id,
                chunk_index=idx,
                audience=rec.audience or "staff",
                text=text_val,
                token_count=_count_tokens_approx(text_val),
                sha256=_sha256_text(text_val),
                start_ms=start_ms,
                end_ms=end_ms,
                page_num=page_num,
                created_at=datetime.utcnow(),
            ))
        db.add_all(new_rows)
        db.commit()
        chunk_rows = new_rows

        # Best-effort paginated preview for office/book-like files.
        if ext in (".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".rtf", ".epub", ".fb2"):
            preview_pdf = _generate_preview_pdf(rec.storage_path)
            if preview_pdf:
                _assign_chunk_pages_from_preview(chunk_rows, preview_pdf)
                db.add_all(chunk_rows)
                db.commit()

    settings = get_effective_gigachat_settings(db, rec.portal_id)
    model = (settings.get("embedding_model") or settings.get("model") or "").strip()
    api_base = (settings.get("api_base") or "").strip()
    if not model:
        rec.status = "error"
        rec.error_message = "missing_embedding_model"
        db.add(rec)
        db.commit()
        return {"ok": False, "error": "missing_embedding_model"}
    token, err = get_valid_gigachat_access_token(db)
    if err or not token:
        rec.status = "error"
        rec.error_message = err or "missing_access_token"
        db.add(rec)
        db.commit()
        return {"ok": False, "error": rec.error_message}

    # remove existing embeddings for this model only (keep other models)
    chunk_ids = [c.id for c in chunk_rows]
    if chunk_ids:
        db.execute(delete(KBEmbedding).where(
            KBEmbedding.chunk_id.in_(chunk_ids),
            KBEmbedding.model == model,
        ))
        db.commit()

    vectors: list[list[float]] = []
    usage_tokens_total = 0
    base_batch = 6
    # retry with backoff on 429 (rate limit)
    retry_delays = [2, 4, 8]

    def _unpack_embeddings(result):
        if isinstance(result, tuple):
            if len(result) == 3:
                return result
            if len(result) == 2:
                embeds, err = result
                return embeds, err, None
        return None, "invalid_embeddings_response", None

    for i in range(0, len(chunk_rows), base_batch):
        batch = chunk_rows[i:i + base_batch]
        texts = [c.text for c in batch]
        embeds = None
        err = None
        # try with shrinking batch sizes on retry
        batch_sizes = [len(texts), max(1, len(texts) // 2), 1]
        for attempt, delay in enumerate([0] + retry_delays):
            if delay:
                import time
                time.sleep(delay)
            size = batch_sizes[min(attempt, len(batch_sizes) - 1)]
            subset = texts[:size]
            embeds, err, usage = _unpack_embeddings(create_embeddings(api_base, token, model, subset))
            if err and "401" in err:
                token, err2 = get_valid_gigachat_access_token(db, force_refresh=True)
                if token and not err2:
                    embeds, err, usage = _unpack_embeddings(create_embeddings(api_base, token, model, subset))
            if err and "429" in err:
                continue
            if err:
                break
            if embeds:
                vectors.extend(embeds)
                if isinstance(usage, dict) and usage.get("total_tokens"):
                    usage_tokens_total += int(usage.get("total_tokens") or 0)
                if size < len(texts):
                    # process remaining texts in the batch
                    rest = texts[size:]
                    for j in range(0, len(rest), 1):
                        embeds2, err2, _usage2 = _unpack_embeddings(create_embeddings(api_base, token, model, [rest[j]]))
                        if err2 and "401" in err2:
                            token, err3 = get_valid_gigachat_access_token(db, force_refresh=True)
                            if token and not err3:
                                embeds2, err2, _usage2 = _unpack_embeddings(create_embeddings(api_base, token, model, [rest[j]]))
                        if err2 and "429" in err2:
                            err = "rate_limited"
                            break
                        if err2 or not embeds2:
                            err = err2 or "embedding_failed"
                            break
                        vectors.extend(embeds2)
                        if isinstance(_usage2, dict) and _usage2.get("total_tokens"):
                            usage_tokens_total += int(_usage2.get("total_tokens") or 0)
                break
        if err:
            if "429" in str(err) or err == "rate_limited":
                rec.status = "queued"
                rec.error_message = "rate_limited"
                db.add(rec)
                db.commit()
                return {"ok": False, "error": "rate_limited"}
            rec.status = "error"
            rec.error_message = err or "embedding_failed"
            db.add(rec)
            db.commit()
            return {"ok": False, "error": rec.error_message}

    emb_rows: list[KBEmbedding] = []
    for ch, vec in zip(chunk_rows, vectors):
        emb_rows.append(KBEmbedding(
            chunk_id=ch.id,
            vector_json=vec,
            model=model,
            dim=len(vec) if vec else None,
            created_at=datetime.utcnow(),
        ))
    db.add_all(emb_rows)
    db.flush()
    for emb_row, vec in zip(emb_rows, vectors):
        if emb_row.id:
            write_vector_column(db, int(emb_row.id), vec)
    db.commit()
    # record embedding usage (portal-level, no user)
    try:
        pricing = get_pricing(db)
        usage_tokens = usage_tokens_total if usage_tokens_total > 0 else None
        cost = calc_cost_rub(usage_tokens, pricing.get("embed_rub_per_1k", 0.0))
        record_usage(
            db,
            portal_id=rec.portal_id,
            user_id=None,
            request_id=f"file:{rec.id}",
            kind="embedding",
            model=model,
            tokens_prompt=None,
            tokens_completion=None,
            tokens_total=int(usage_tokens) if usage_tokens else None,
            cost_rub=cost,
            status="ok",
            error_code=None,
        )
    except Exception:
        pass
    rec.status = "ready"
    rec.error_message = None
    rec.processed_at = datetime.utcnow()
    db.add(rec)
    db.commit()
    return {"ok": True, "chunks": len(chunk_rows)}
